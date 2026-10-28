"""Build the single Word report from reviewed content; no ML imports or simulation.

Use the bundled document runtime, after the artifact-operation marker. Run via
the installed Windows guard when agent-operated. Render and inspect all pages
with the Documents skill before delivering the generated DOCX.
"""
import argparse
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def build_report(output):
    output = Path(output)
    if output.exists():
        raise FileExistsError(f'Preserve the existing report before editing: {output}')
    content = json.loads((ROOT / 'docs/report_content.json').read_text(encoding='utf-8'))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, size in [('Normal', 11), ('Title', 22), ('Subtitle', 11),
                       ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.widow_control = True
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Calibri')
        if name.startswith('Heading'):
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(12)
    # No title rule, theme color, or decorative graphics.
    for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        props = doc.styles[name].element.find(qn('w:pPr'))
        if props is not None:
            for border in list(props.findall(qn('w:pBdr'))):
                props.remove(border)
    doc.add_paragraph(content['title'], 'Title')
    doc.add_paragraph(content['subtitle'], 'Subtitle')
    doc.add_paragraph(content['date'])

    for index, item in enumerate(content['sections'], 1):
        doc.add_heading(f"{index} {item['title']}", level=1)
        for paragraph in item['paragraphs']:
            doc.add_paragraph(paragraph)
        # These are executable-source expressions, not LaTeX equations.
        for line in item.get('code', []):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        for data in item.get('tables', []):
            table = doc.add_table(rows=1, cols=len(data['headers']))
            table.style = 'Table Grid'
            table.autofit = False
            widths = [16.6 / len(data['headers'])] * len(data['headers'])
            if len(widths) == 2:
                widths = [6.3, 10.3]
            for col, width in zip(table.columns, widths):
                col.width = Cm(width)
            for cell, text in zip(table.rows[0].cells, data['headers']):
                cell.text = text
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in data['rows']:
                for cell, text in zip(table.add_row().cells, values):
                    cell.text = str(text)
            for row in table.rows:
                no_split = OxmlElement('w:cantSplit')
                row._tr.get_or_add_trPr().append(no_split)
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.space_before = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(10)
            doc.add_paragraph()

    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Trang ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = content['title']
    doc.core_properties.subject = 'Mario PPO source implementation and experimental results'
    doc.core_properties.author = ''
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    with zipfile.ZipFile(output) as archive:
        if archive.testzip() is not None or any(n.startswith('word/media/') for n in archive.namelist()):
            raise ValueError('DOCX integrity or no-images requirement failed')
    print(json.dumps({'output': str(output), 'sections': len(content['sections']),
                      'visual_review': 'pending', 'images': 0}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', type=Path, default=ROOT / 'docs/BAO_CAO_MARIO.docx')
    build_report(parser.parse_args().out)
